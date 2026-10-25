#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把合并版 Markdown 转成 docx（他们给的格式），供 Ann 直接发回去。

用 python-docx 手写转换：标题/表格/列表/代码块/引用。
不用 pandoc（本机没装）。表格是这份文档的主体，所以表格必须单独处理。
"""
import re, sys, os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = sys.argv[1] if len(sys.argv) > 1 else 'MIRU-规范与现状-合并版.md'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'MIRU-规范与现状-合并版.docx'

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(10)

def add_para(text, bold=False, italic=False, size=None, color=None, style_name=None):
    p = doc.add_paragraph(style=style_name) if style_name else doc.add_paragraph()
    # 行内 **粗体** 与 `代码`
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else:
            r = p.add_run(part)
        r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    return p

lines = open(SRC, encoding='utf-8').read().split('\n')
i = 0
in_code = False
code_buf = []

while i < len(lines):
    ln = lines[i]

    # 代码块
    if ln.strip().startswith('```'):
        if in_code:
            for c in code_buf:
                p = doc.add_paragraph()
                r = p.add_run(c if c else ' ')
                r.font.name = 'Consolas'; r.font.size = Pt(9)
            code_buf = []; in_code = False
        else:
            in_code = True
        i += 1
        continue
    if in_code:
        code_buf.append(ln)
        i += 1
        continue

    # 表格：连续的 | 开头的行
    if ln.strip().startswith('|') and i + 1 < len(lines) and re.match(r'^\s*\|[\s\-:|]+\|\s*$', lines[i+1]):
        header = [c.strip() for c in ln.strip().strip('|').split('|')]
        i += 2
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
            i += 1
        t = doc.add_table(rows=1, cols=len(header))
        t.style = 'Light Grid Accent 1'
        for j, h in enumerate(header):
            cell = t.rows[0].cells[j]
            cell.text = ''
            r = cell.paragraphs[0].add_run(h.replace('**', ''))
            r.bold = True; r.font.size = Pt(9)
        for row in rows:
            cells = t.add_row().cells
            for j in range(len(header)):
                val = row[j] if j < len(row) else ''
                cells[j].text = ''
                r = cells[j].paragraphs[0].add_run(val.replace('**', ''))
                r.font.size = Pt(9)
        doc.add_paragraph()
        continue

    # 标题
    m = re.match(r'^(#{1,4})\s+(.*)$', ln)
    if m:
        level = len(m.group(1))
        h = doc.add_heading(level=min(level, 4))
        run = h.add_run(m.group(2).strip())
        run.font.name = 'Microsoft YaHei'
        i += 1
        continue

    # 引用
    if ln.strip().startswith('>'):
        add_para(ln.strip().lstrip('>').strip(), italic=True, size=9,
                 color=RGBColor(0x55, 0x55, 0x55))
        i += 1
        continue

    # 分隔线
    if re.match(r'^\s*---+\s*$', ln):
        doc.add_paragraph('─' * 40)
        i += 1
        continue

    # 列表
    m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', ln)
    if m:
        add_para(m.group(3), style_name='List Bullet' if m.group(2) in '-*' else 'List Number')
        i += 1
        continue

    if ln.strip() == '':
        i += 1
        continue

    add_para(ln.strip())
    i += 1

doc.save(OUT)
print('已写', OUT, os.path.getsize(OUT), 'bytes')
