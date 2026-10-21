import json, os, re, sys
sys.path.insert(0, os.path.dirname(__file__))
from assets94 import TSUNS, CROWN, OFFICE, STYLE_BG, STYLE_PT, STYLE_CG, plan
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
H = os.path.dirname(os.path.abspath(__file__))
P = plan()
rows = []  # story, type, file, 用途, 中文说明, prompt, 尺寸, 优先级
def add(st, ty, fn, use, zh, pr, size, pri): rows.append([st, ty, fn, use, zh, pr, size, pri])
STORY = {'tsuns': '王不见王', 'crown': '王冠的最后誓言', 'office': '冰场之外'}
# covers
add('tsuns', '封面', 'tsuns_cover.jpg', '发现页 / 故事页封面', '海报：两个名字一上一下，男女背对背站在闪光灯里', 'movie poster style, a Chinese actor (33, suit buttoned at the second button, old wristwatch) and a Chinese actress (30, grey hoodie, red pen) standing back to back under camera flashes, two names on the poster one above the other, moody teal and amber, elegant title space at top, ' + 'vertical 3:4, no text', '3:4（1080×1440）', '必需')
add('crown', '封面', 'crown_cover.jpg', '发现页 / 故事页封面', '加冕大殿，王储不跪，向侧台的记录员伸出手', 'romantasy book cover, a young crown prince (26, dark hair, formal regalia, no crown on his head) refusing to kneel at a coronation altar, extending his hand toward a female royal scribe holding parchment at the side stage, candlelight, stained glass, golden crown glowing, vertical 3:4, no text', '3:4（1080×1440）', '必需')
add('office', '封面', 'office_cover.jpg', '发现页 / 故事页封面', '冰场边，队长背影，手机屏上停住的录像', 'sports romance cover, college ice hockey captain (22, short brown hair, worn training jacket) standing at the rink boards at night, looking back over his shoulder, a phone screen with a paused game video in the foreground, cold blue rink light, vertical 3:4, no text', '3:4（1080×1440）', '必需')
# portraits
PT = {
 'tsuns': [('lu', '陆识野', 'Chinese male film star and producer, 33, sharp restrained face, short black hair, dark suit with only the second button fastened, old wristwatch, faint scar of an old left-shoulder injury hidden under the suit, calm guarded gaze'),
           ('fang', '方叙', 'Chinese film producer, 40, neat glasses, phone pressed to his ear, holding three folders, tailored shirt with rolled sleeves, shrewd but not cruel'),
           ('xiaoman', '林小满', 'Chinese talent agent, 28, ponytail, carrying a row of high heels in one hand, phone in the other, tired but warm eyes'),
           ('zhou', '老周', 'Chinese film director, 55, baseball cap, reading glasses pushed up, grey stubble, sitting behind a monitor, knowing half-smile')],
 'crown': [('caelan', '凯兰', 'fantasy crown prince, 26, dark wavy hair, formal navy and gold regalia without a crown, thin leather clock-repair gloves, steady kind but burdened eyes'),
           ('rowan', '罗温', 'royal guard commander, 27, short cropped hair, silver-and-blue plate armor, a flour smudge on one gauntlet, upright and loyal'),
           ('elira', '埃莉拉', 'commoner council representative, 30, practical wool coat, hair tied back, a bundle of petitions under her arm, sharp determined eyes'),
           ('marek', '马雷克', 'regent advisor, 51, grey trimmed beard, dark velvet robes, ledger in hand, cold calculating stare')],
 'office': [('evan', '埃文', 'college ice hockey captain, 22, short brown hair, old team training jacket, tape on his stick hand, serious direct gaze, slight tiredness'),
            ('mateo', '马特奥', 'college hockey goalie, 21, curly dark hair, holding a vintage film camera, easy grin, goalie pads on his shoulder'),
            ('dana', '丹娜', 'ice hockey head coach, 45, woman, short hair, team jacket, clipboard and whistle, firm fair expression'),
            ('harper', '哈珀', 'team doctor, 30, woman, medical polo shirt, stethoscope, tablet in hand, calm professional'),
            ('mason', '梅森', 'college hockey player, 21, buzz cut, competitive smirk, practice jersey, arms crossed')],
}
for st, L in PT.items():
    for key, zh, pr in L: add(st, '人物立绘', '%s_char_%s.jpg' % (st, key), '角色头像 / 聊天头像 / 来电头像', zh, pr + ', ' + STYLE_PT, '3:4（900×1200）', '必需')
# backgrounds
for st, T in (('tsuns', TSUNS), ('crown', CROWN), ('office', OFFICE)):
    from collections import Counter
    cnt = Counter(P[st].values())
    for cid, words, zh, pr in T:
        n = cnt.get(cid, 0)
        if not n: continue
        add(st, '场景背景', '%s_%s.jpg' % (st, cid), '用于 %d 个场景' % n, zh, pr + ', ' + STYLE_BG, '9:16（1080×1920）', '必需' if n >= 3 else '建议')
# key CGs
CG_T = [(1, 4, '海报拍摄棚：两人站位，谁也不退半步', 'photo studio poster shoot, actor and actress standing on the same level refusing to step back, photographer adjusting lights'),
        (2, 3, '雨棚下第六条，导演没喊卡', 'film set rain shelter, actor and actress facing each other three meters apart under rain machines, director frozen behind monitor'),
        (3, 2, '发布会台上被问"谁在上面"', 'press conference stage, actress at the microphone, actor beside her looking away, flashes everywhere'),
        (4, 4, '杀青那一镜：第三条，不是演的', 'last take of the film, rainy street set, the two leads close, his forehead almost touching hers, crew silent'),
        (5, 5, '天台：两下敲击', 'night rooftop, he knocks twice on the railing, she stands a few steps away, city lights'),
        (6, 4, '桌下那只手', 'long press table, under the table his hand rests next to hers, faces calm above'),
        (7, 7, '冷雨夜', 'cold rain at night, actress alone under a streetlight, umbrella closed in her hand'),
        (8, 7, '天台烟花夜戏', 'rooftop fireworks scene, the two leads on opposite ends of the rooftop, fireworks between them'),
        (9, 10, '他第一次喊她全名', 'apartment doorway at night, he says her full name, she turns'),
        (10, 3, '首映礼红毯，双日同天', 'premiere red carpet, the two leads arriving at the same moment, crowd screaming'),
        (11, 10, '机场安检口：谁让一步', 'airport security gate at dawn, he holds a suitcase, she stands behind the yellow line'),
        (12, 6, '结局A：两个名字并排', 'giant poster wall, two names side by side, the pair standing shoulder to shoulder'),
        (12, 7, '结局B：两台摄影机对着彼此', 'unnamed studio, two cameras facing each other, both of them behind the lenses'),
        (12, 8, '结局C：十年后的酒馆', 'old bar ten years later, older actor waiting, the actress walks in')]
for ch, sc, zh, pr in CG_T: add('tsuns', '剧情CG（可选）', 'tsuns_cg_ch%02d_%02d.jpg' % (ch, sc), '第%d章关键时刻' % ch, zh, pr + ', ' + STYLE_CG, '9:16（1080×1920）', '可选')
for st in ('crown', 'office'):
    meta = json.load(open(os.path.join(H, st + '.md.meta.json')))
    last = {}
    for b in meta['bg']: last[b['ch']] = b
    who = '凯兰（王储）和你（女记录员）' if st == 'crown' else '埃文（队长）和你（新队员）'
    wen = 'a young crown prince and a female royal scribe' if st == 'crown' else 'a college hockey captain and a new teammate'
    for ch in sorted(last):
        b = last[ch]
        add(st, '剧情CG（可选）', '%s_cg_ch%02d.jpg' % (st, ch), '第%d章收尾时刻' % ch, who + '：' + b['desc'][:60], '（中文提示词）电影剧照，竖屏9:16，' + who + '同框，成年人，情绪关键时刻，浅景深，无文字。场景：' + b['desc'][:160], '9:16（1080×1920）', '可选')
# videos
VID = [('tsuns', 'vid_open', '开场：快门声，戴旧表的手按住海报上第二个名字', 'camera shutter clicks seven times, a hand with an old wristwatch presses a movie poster proof onto a table, fingertip on the second name, slow push-in, moody studio light'),
       ('tsuns', 'vid_final', '终章：海报墙的灯一盏盏亮起，两个名字并排', 'lights switch on one by one along a giant poster wall, two names side by side at the same height, two people step into frame shoulder to shoulder'),
       ('crown', 'vid_open', '开场：加冕大典，王储拒绝下跪，转身向侧台伸手', 'coronation hall full of candles, the crown prince stops before the altar, does not kneel, turns and extends his hand toward the side stage, nobles gasp'),
       ('crown', 'vid_final', '尾声：春日塔楼露台，花瓣吹过王城', 'spring tower balcony, petals blowing across the city, the prince and the scribe look over the kingdom, slow pull-back'),
       ('office', 'vid_open', '开场：更衣室，手机里的录像停在碰撞前一帧', 'hockey locker room, a phone shows a paused clip one frame before a collision, teammates turn, the captain steps forward, cold white light'),
       ('office', 'vid_final', '终章：夏日湖边，队长回头', 'summer campus lake at sunset, the hockey captain sits on a bench with his stick, turns and smiles softly, warm light')]
for st, slot, zh, pr in VID: add(st, '视频', '%s_%s.mp4' % (st, slot), '开场 / 终章自动播放（静音循环）', zh, pr + ', vertical 9:16, 5 seconds, cinematic', '9:16 · 5秒 · 720p', '建议')
wb = Workbook(); ws = wb.active; ws.title = '全部素材'
hdr = ['剧本', '类型', '文件名（请按这个命名）', '用在哪', '画面（中文）', '提示词（英文，可直接用）', '尺寸', '优先级']
ws.append(hdr)
for r in rows: ws.append([STORY[r[0]]] + r[1:])
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='5B3A6B')
for col, w in zip('ABCDEFGH', [14, 12, 30, 22, 40, 90, 18, 8]): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'A2'
s2 = wb.create_sheet('数量汇总')
s2.append(['剧本', '封面', '人物立绘', '场景背景', '剧情CG（可选）', '视频'])
for st in STORY:
    f = lambda t: sum(1 for r in rows if r[0] == st and r[1] == t)
    s2.append([STORY[st], f('封面'), f('人物立绘'), f('场景背景'), f('剧情CG（可选）'), f('视频')])
s2.append(['合计'] + [sum(1 for r in rows if r[1] == t) for t in ['封面', '人物立绘', '场景背景', '剧情CG（可选）', '视频']])
for c in s2[1]: c.font = Font(bold=True)
s2.column_dimensions['A'].width = 18
out = '/mnt/user-data/outputs/MIRU_三部新剧_图片视频提示词.xlsx'
wb.save(out)
json.dump({'rows': rows}, open(os.path.join(H, 'prompts.json'), 'w'), ensure_ascii=False)
print(out, len(rows), [list(r) for r in s2.iter_rows(values_only=True)])
